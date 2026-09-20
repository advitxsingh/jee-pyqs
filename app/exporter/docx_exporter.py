"""
Word Document (.docx) Exporter for JEE PYQ Concept Sheets & Revision Notes.
Generates a print-ready, professionally styled revision booklet.
"""

import os
import re
from typing import Optional, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.db.database import (
    get_chapter, get_chapter_synthesis, get_concepts_by_chapter,
    get_questions_by_chapter, get_question_by_id, DB_PATH
)

# Palette
COLOR_PRIMARY = RGBColor(26, 54, 93)      # Deep Navy
COLOR_SECONDARY = RGBColor(43, 108, 176)  # Slate Blue
COLOR_ACCENT = RGBColor(197, 48, 48)      # Burgundy / Warning
COLOR_MUTED = RGBColor(113, 128, 150)     # Gray
COLOR_DARK = RGBColor(45, 55, 72)         # Off-black


def format_math_for_print(text: str) -> str:
    """
    Format LaTeX expressions into clean, legible typographic notation for print documents.
    """
    if not text:
        return ""

    def _clean_formula(tex: str) -> str:
        s = tex.strip()
        s = re.sub(r'\\begin\{(?:aligned|matrix|array|cases)\}', '', s)
        s = re.sub(r'\\end\{(?:aligned|matrix|array|cases)\}', '', s)
        s = s.replace('&', ' ').replace(r'\\', '\n')
        s = re.sub(r'\\text\{([^}]*)\}', r'\1', s)
        s = re.sub(r'\\mathrm\{([^}]*)\}', r'\1', s)
        s = re.sub(r'\\mathbf\{([^}]*)\}', r'\1', s)

        for _ in range(3):
            s = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1 / \2)', s)

        s = re.sub(r'\\sqrt\[([^]]+)\]\{([^}]+)\}', r'\1√(\2)', s)
        s = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', s)
        s = re.sub(r'\\sqrt\s*([a-zA-Z0-9])', r'√\1', s)

        s = re.sub(r'(\\(?:alpha|beta|gamma|delta|epsilon|zeta|eta|theta|iota|kappa|lambda|mu|nu|xi|pi|rho|sigma|tau|upsilon|phi|chi|psi|omega|Gamma|Delta|Theta|Lambda|Xi|Pi|Sigma|Upsilon|Phi|Psi|Omega))([a-zA-Z0-9\\])', r'\1 \2', s)
        s = re.sub(r'\\vec\{([^}]+)\}', r'vec(\1)', s)
        s = re.sub(r'\\hat\{([^}]+)\}', r'\1̂', s)
        s = re.sub(r'\\bar\{([^}]+)\}', r'bar(\1)', s)
        s = re.sub(r'\\overset\{[^}]*\}\{([^}]+)\}', r'vec(\1)', s)

        s = s.replace(r'^{\circ}', '°').replace(r'^\circ', '°')
        s = re.sub(r'\^\{([0-9\+\-]+)\}', r'^\1', s)
        s = re.sub(r'\^([0-9])', r'^\1', s)
        s = re.sub(r'\_\{([^}]+)\}', r'_(\1)', s)

        GREEK_MAP = {
            r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ', r'\delta': 'δ',
            r'\epsilon': 'ε', r'\zeta': 'ζ', r'\eta': 'η', r'\theta': 'θ',
            r'\iota': 'ι', r'\kappa': 'κ', r'\lambda': 'λ', r'\mu': 'μ',
            r'\nu': 'ν', r'\xi': 'ξ', r'\pi': 'π', r'\rho': 'ρ',
            r'\sigma': 'σ', r'\tau': 'τ', r'\upsilon': 'υ', r'\phi': 'ϕ',
            r'\chi': 'χ', r'\psi': 'ψ', r'\omega': 'ω',
            r'\Gamma': 'Γ', r'\Delta': 'Δ', r'\Theta': 'Θ', r'\Lambda': 'Λ',
            r'\Xi': 'Ξ', r'\Pi': 'Π', r'\Sigma': 'Σ', r'\Upsilon': 'Υ',
            r'\Phi': 'Φ', r'\Psi': 'Ψ', r'\Omega': 'Ω',
        }
        for k, v in GREEK_MAP.items():
            s = re.sub(re.escape(k) + r'(?![a-zA-Z])', v, s)

        SYM_MAP = {
            r'\times': '×', r'\cdot': '·', r'\div': '÷', r'\pm': '±', r'\mp': '∓',
            r'\longrightarrow': '→', r'\rightarrow': '→', r'\leftarrow': '←',
            r'\Rightarrow': '⇒', r'\rightleftharpoons': '⇌',
            r'\le': '≤', r'\ge': '≥', r'\ne': '≠', r'\approx': '≈',
            r'\equiv': '≡', r'\in': '∈', r'\perp': '⊥', r'\parallel': '∥',
            r'\infty': '∞', r'\int': '∫', r'\sum': '∑', r'\partial': '∂',
            r'\nabla': '∇', r'\ell': 'ℓ', r'\therefore': '∴', r'\because': '∵',
            r'\dots': '...', r'\cdots': '...',
            r'\log': 'log', r'\ln': 'ln', r'\sin': 'sin', r'\cos': 'cos', r'\tan': 'tan',
            r'\quad': ' ', r'\qquad': '  ', r'\,': ' ', r'\;': ' ', r'\!': ''
        }
        for k, v in SYM_MAP.items():
            s = re.sub(re.escape(k) + r'(?![a-zA-Z])', v, s)

        s = s.replace('{', '').replace('}', '')
        return re.sub(r'\s+', ' ', s).strip()

    res = re.sub(r'\$\$([^\$]+)\$\$', lambda m: _clean_formula(m.group(1)), text)
    res = re.sub(r'\$([^\$]+)\$', lambda m: _clean_formula(m.group(1)), res)
    return res.replace('\u2061', '').replace('\u200b', '')


def set_cell_background(cell, fill_color_hex: str):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def generate_chapter_docx(chapter_slug: str, output_path: Optional[str] = None, db_path: str = DB_PATH) -> str:
    """
    Generate a formatted Word document for the specified chapter.
    Returns the file path of the generated .docx file.
    """
    chapter = get_chapter(chapter_slug, db_path=db_path)
    if not chapter:
        raise ValueError(f"Chapter '{chapter_slug}' not found.")

    synthesis = get_chapter_synthesis(chapter_slug, db_path=db_path)
    concepts = get_concepts_by_chapter(chapter_slug, db_path=db_path)
    questions = get_questions_by_chapter(chapter_slug, limit=1000, db_path=db_path)

    if not output_path:
        os.makedirs("data/exports", exist_ok=True)
        output_path = f"data/exports/{chapter_slug}_revision_notes.docx"

    doc = Document()

    # Configure Margins (0.75 in)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Document Title / Cover Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run(f"JEE MAIN REVISION MASTER: {chapter.title.upper()}")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run(f"Subject: {chapter.subject.title()}  |  Target: JEE Main  |  PYQs Analyzed: {len(questions)}")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_MUTED

    # Section 1: Executive Blueprint Table
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. High-Yield Concept Priority Matrix")
    r_h1.font.color.rgb = COLOR_PRIMARY
    r_h1.font.size = Pt(14)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Priority Tier", "Concept Name", "Frequency", "Category"]
    widths = [Inches(1.5), Inches(3.0), Inches(1.2), Inches(1.3)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1A365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Order concepts: High/Very High curated concepts first, fallback at end
    def _concept_sort_key(c):
        is_fallback = c.name.endswith("Core Application")
        tier_weight = {"Very High": 4, "High": 3, "Moderate": 2, "Standard": 1}.get(c.frequency_tier, 0)
        return (1 if is_fallback else 0, -tier_weight, -c.exam_frequency)

    sorted_concepts = sorted(concepts, key=_concept_sort_key)

    for idx, c in enumerate(sorted_concepts):
        row_cells = table.add_row().cells
        bg_color = "F7FAFC" if idx % 2 == 0 else "FFFFFF"
        for i, w in enumerate(widths):
            row_cells[i].width = w
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=80, bottom=80, left=150, right=150)

        # Tier
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(c.frequency_tier)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        if c.frequency_tier == "Very High":
            r0.font.color.rgb = COLOR_ACCENT
        elif c.frequency_tier == "High":
            r0.font.color.rgb = COLOR_SECONDARY
        else:
            r0.font.color.rgb = COLOR_MUTED

        # Name
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(c.name)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

        # Frequency
        p2 = row_cells[2].paragraphs[0]
        r2 = p2.add_run(f"{c.exam_frequency} PYQs")
        r2.font.size = Pt(9.5)

        # Category
        p3 = row_cells[3].paragraphs[0]
        r3 = p3.add_run(c.category)
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 2: Master Formula Sheet
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Master Formula Reference Card")
    r_h2.font.color.rgb = COLOR_PRIMARY
    r_h2.font.size = Pt(14)

    f_table = doc.add_table(rows=1, cols=2)
    f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_headers = ["Concept / Topic", "Mathematical Formulation & Standard Equations"]
    f_widths = [Inches(2.5), Inches(4.5)]

    f_hdr = f_table.rows[0].cells
    for i, title in enumerate(f_headers):
        f_hdr[i].width = f_widths[i]
        set_cell_background(f_hdr[i], "2B6CB0")
        set_cell_margins(f_hdr[i], top=120, bottom=120, left=150, right=150)
        p = f_hdr[i].paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    formula_idx = 0
    for c in sorted_concepts:
        if not c.standard_formulas or c.name.endswith("Core Application"):
            continue
        row_cells = f_table.add_row().cells
        bg_color = "F7FAFC" if formula_idx % 2 == 0 else "FFFFFF"
        formula_idx += 1
        for i, w in enumerate(f_widths):
            row_cells[i].width = w
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=150, right=150)

        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(c.name)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        clean_formula_text = format_math_for_print(c.standard_formulas)
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(clean_formula_text)
        r1.font.name = "Cambria Math"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 3: Deep-Dive Concepts with Linked PYQs
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Concept Deep-Dive & Source PYQ Evidence")
    r_h3.font.color.rgb = COLOR_PRIMARY
    r_h3.font.size = Pt(14)

    for idx, c in enumerate(sorted_concepts, 1):
        ch_title = doc.add_heading(level=2)
        r_c = ch_title.add_run(f"3.{idx} {c.name} ({c.frequency_tier} Yield — {c.exam_frequency} PYQs)")
        r_c.font.color.rgb = COLOR_SECONDARY
        r_c.font.size = Pt(12)

        # Summary
        p_sum = doc.add_paragraph()
        r_s_lbl = p_sum.add_run("Core Principle: ")
        r_s_lbl.font.bold = True
        r_s_lbl.font.size = Pt(10)
        r_s_val = p_sum.add_run(format_math_for_print(c.summary))
        r_s_val.font.size = Pt(10)

        # Trap Callout Box
        if c.common_traps:
            callout = doc.add_table(rows=1, cols=1)
            callout.alignment = WD_TABLE_ALIGNMENT.CENTER
            c_cell = callout.rows[0].cells[0]
            c_cell.width = Inches(7.0)
            set_cell_background(c_cell, "FFF5F5")  # Soft red
            set_cell_margins(c_cell, top=100, bottom=100, left=150, right=150)
            cp = c_cell.paragraphs[0]
            r_tr_lbl = cp.add_run("⚠️ Common Trap / JEE Pitfall: ")
            r_tr_lbl.font.bold = True
            r_tr_lbl.font.size = Pt(9.5)
            r_tr_lbl.font.color.rgb = COLOR_ACCENT
            r_tr_val = cp.add_run(format_math_for_print(c.common_traps))
            r_tr_val.font.size = Pt(9.5)
            r_tr_val.font.color.rgb = COLOR_DARK

        # Speed Trick Callout Box
        if c.tips_and_tricks:
            trick_box = doc.add_table(rows=1, cols=1)
            trick_box.alignment = WD_TABLE_ALIGNMENT.CENTER
            t_cell = trick_box.rows[0].cells[0]
            t_cell.width = Inches(7.0)
            set_cell_background(t_cell, "F0FFF4")  # Soft green
            set_cell_margins(t_cell, top=100, bottom=100, left=150, right=150)
            tp = t_cell.paragraphs[0]
            r_tk_lbl = tp.add_run("💡 Speed Hack / Trick: ")
            r_tk_lbl.font.bold = True
            r_tk_lbl.font.size = Pt(9.5)
            r_tk_lbl.font.color.rgb = RGBColor(39, 103, 73)
            r_tk_val = tp.add_run(format_math_for_print(c.tips_and_tricks))
            r_tk_val.font.size = Pt(9.5)
            r_tk_val.font.color.rgb = COLOR_DARK

        # Linked PYQs
        pyq_hdr = doc.add_paragraph()
        pyq_hdr.paragraph_format.space_before = Pt(6)
        pyq_hdr.paragraph_format.space_after = Pt(2)
        r_ph = pyq_hdr.add_run("Source PYQs testing this concept:")
        r_ph.font.bold = True
        r_ph.font.size = Pt(9.5)
        r_ph.font.color.rgb = COLOR_PRIMARY

        matched_qs = [q for q in questions if c.name in q.key_concepts or q.topic_tag == c.name][:3]
        if not matched_qs:
            matched_qs = questions[:2]

        for mq in matched_qs:
            qp = doc.add_paragraph(style='List Bullet')
            qp.paragraph_format.space_after = Pt(2)
            rq1 = qp.add_run(f"[{mq.year} {mq.shift or 'Shift'}] Q#{mq.question_index} ({mq.question_type}): ")
            rq1.font.bold = True
            rq1.font.size = Pt(9)
            clean_q_snippet = format_math_for_print(mq.question_text[:140])
            rq2 = qp.add_run(f"\"{clean_q_snippet}...\"  ")
            rq2.font.size = Pt(9)
            rq3 = qp.add_run(f"[Ans: {mq.correct_answer or 'N/A'}]")
            rq3.font.bold = True
            rq3.font.size = Pt(9)
            rq3.font.color.rgb = COLOR_SECONDARY

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 4: Benchmark Illustrative PYQs
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Benchmark Illustrative PYQs & Master Solutions")
    r_h4.font.color.rgb = COLOR_PRIMARY
    r_h4.font.size = Pt(14)

    benchmark_ids = synthesis.benchmark_pyq_ids if synthesis else []
    benchmarks = []
    for bid in benchmark_ids:
        b = get_question_by_id(str(bid), db_path=db_path)
        if b:
            benchmarks.append(b)

    if not benchmarks:
        benchmarks = [q for q in questions if len(q.explanation_text or "") > 150][:4]

    for b_idx, bq in enumerate(benchmarks, 1):
        bp_head = doc.add_paragraph()
        bp_head.paragraph_format.space_before = Pt(8)
        bp_head.paragraph_format.space_after = Pt(2)
        r_bh = bp_head.add_run(f"Benchmark Problem {b_idx}: {bq.paper_name} (Q#{bq.question_index})")
        r_bh.font.bold = True
        r_bh.font.size = Pt(11)
        r_bh.font.color.rgb = COLOR_PRIMARY

        # Question Statement
        bp_q = doc.add_paragraph()
        bp_q.paragraph_format.space_after = Pt(4)
        rq = bp_q.add_run(format_math_for_print(bq.question_text))
        rq.font.size = Pt(10)

        # Options if MCQ
        if bq.options:
            for opt in bq.options:
                op = doc.add_paragraph(style='List Bullet')
                op.paragraph_format.space_after = Pt(1)
                r_opt = op.add_run(f"({opt.label}) {format_math_for_print(opt.text)}")
                r_opt.font.size = Pt(9.5)
                if opt.label == bq.correct_answer:
                    r_opt.font.bold = True
                    r_opt.font.color.rgb = RGBColor(39, 103, 73)

        # Solution Box
        sol_box = doc.add_table(rows=1, cols=1)
        sol_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        sb_cell = sol_box.rows[0].cells[0]
        sb_cell.width = Inches(7.0)
        set_cell_background(sb_cell, "F7FAFC")
        set_cell_margins(sb_cell, top=100, bottom=100, left=150, right=150)
        
        sp = sb_cell.paragraphs[0]
        r_sa = sp.add_run(f"Correct Answer: {bq.correct_answer}\n")
        r_sa.font.bold = True
        r_sa.font.size = Pt(9.5)
        r_sa.font.color.rgb = COLOR_PRIMARY

        clean_expl = format_math_for_print(bq.explanation_text or "")
        r_se = sp.add_run(f"Solution Analysis:\n{clean_expl}")
        r_se.font.size = Pt(9)
        r_se.font.color.rgb = COLOR_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_path)
    return output_path
